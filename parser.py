
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract readable text from a PDF or DOCX resume."""

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Upload a PDF or DOCX file."
        )

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))

        if reader.is_encrypted:
            raise ValueError(
                "This PDF is encrypted. Upload an unlocked PDF."
            )

        text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    else:
        document = Document(BytesIO(file_bytes))
        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Also extract text from tables, which some resumes use.
        table_text = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            if cell.text.strip()
        ]

        text = "\n".join(paragraphs + table_text)

    if not text.strip():
        raise ValueError(
            "No readable text found. The resume may be scanned "
            "or image-based. OCR support is not included yet."
        )

    return text
